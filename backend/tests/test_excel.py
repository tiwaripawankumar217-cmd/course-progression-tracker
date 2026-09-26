import io
import pytest
from fastapi.testclient import TestClient
import openpyxl

from backend.main import app
from backend.services.excel_service import (
    excel_service,
    PRIMARY_SHEET_NAME,
    EVIDENCE_SHEET_NAME,
    REQUIRED_COLUMNS,
)
from backend.services.progress_engine import progress_engine
from backend.models.curriculum import Curriculum, LecturePlan
from backend.models.analysis import (
    LectureAnalysisResult,
    TopicTaught,
    ClassworkPerformed,
    HomeworkAssigned,
    MatchedLecture,
)


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def sample_progression_payload():
    return {
        "lecture": "LEC-1",
        "date_taught": "2026-09-26",
        "status": "Ongoing",
        "percent_completed": 66.67,
        "percent_covered": 66.67,
        "classwork": "Find Maximum Element",
        "homework": "Solve 5 Array Problems",
    }


def test_create_workbook(sample_progression_payload):
    """Test 1 — Create workbook: Validates file creation, columns, and cell values."""
    xlsx_bytes = excel_service.generate_progression_workbook(sample_progression_payload)
    assert xlsx_bytes is not None
    assert len(xlsx_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert PRIMARY_SHEET_NAME in wb.sheetnames

    ws = wb[PRIMARY_SHEET_NAME]
    # Verify required columns in row 1
    actual_headers = [ws.cell(row=1, column=c).value for c in range(1, 7)]
    assert actual_headers == REQUIRED_COLUMNS

    # Verify data in row 2
    assert ws.cell(row=2, column=1).value == "2026-09-26"
    assert ws.cell(row=2, column=2).value == "Ongoing"
    assert float(ws.cell(row=2, column=3).value) == 66.67
    assert float(ws.cell(row=2, column=4).value) == 66.67
    assert ws.cell(row=2, column=5).value == "Find Maximum Element"
    assert ws.cell(row=2, column=6).value == "Solve 5 Array Problems"


def test_existing_workbook_update():
    """Test 2 — Existing workbook update: Preserves previous records and appends new one."""
    rec1 = {
        "lecture": "LEC-1",
        "date_taught": "2026-09-20",
        "status": "Completed",
        "percent_completed": 100.0,
        "percent_covered": 100.0,
        "classwork": "Arrays",
        "homework": "Solve 5 problems",
    }
    initial_bytes = excel_service.generate_progression_workbook(rec1)

    rec2 = {
        "lecture": "LEC-2",
        "date_taught": "2026-09-22",
        "status": "Ongoing",
        "percent_completed": 60.0,
        "percent_covered": 75.0,
        "classwork": "Loops",
        "homework": "Practice loops",
    }
    updated_bytes = excel_service.update_progression_workbook(initial_bytes, rec2)

    wb = openpyxl.load_workbook(io.BytesIO(updated_bytes))
    ws = wb[PRIMARY_SHEET_NAME]

    # Row 1 is header, Row 2 is rec1, Row 3 is rec2
    assert ws.cell(row=2, column=1).value == "2026-09-20"
    assert ws.cell(row=2, column=2).value == "Completed"
    assert ws.cell(row=3, column=1).value == "2026-09-22"
    assert ws.cell(row=3, column=2).value == "Ongoing"
    assert ws.max_row == 3


def test_duplicate_update():
    """Test 3 — Duplicate update: Updates existing row instead of appending duplicate."""
    rec1 = {
        "lecture": "LEC-1",
        "date_taught": "2026-09-26",
        "status": "Ongoing",
        "percent_completed": 50.0,
        "percent_covered": 50.0,
        "classwork": "Linear search",
        "homework": "Homework 1",
    }
    initial_bytes = excel_service.generate_progression_workbook(rec1)

    # Re-export same date and lecture with updated metrics
    rec1_updated = {
        "lecture": "LEC-1",
        "date_taught": "2026-09-26",
        "status": "Completed",
        "percent_completed": 100.0,
        "percent_covered": 100.0,
        "classwork": "Linear and Binary search",
        "homework": "Homework 1 completed",
    }
    updated_bytes = excel_service.update_progression_workbook(initial_bytes, rec1_updated)

    wb = openpyxl.load_workbook(io.BytesIO(updated_bytes))
    ws = wb[PRIMARY_SHEET_NAME]

    # Should still only have 1 data row (row 2), updated in place
    assert ws.max_row == 2
    assert ws.cell(row=2, column=1).value == "2026-09-26"
    assert ws.cell(row=2, column=2).value == "Completed"
    assert float(ws.cell(row=2, column=3).value) == 100.0
    assert ws.cell(row=2, column=5).value == "Linear and Binary search"


def test_missing_homework():
    """Test 4 — Missing homework: Missing homework does not crash and leaves empty cell."""
    rec = {
        "lecture": "LEC-1",
        "date_taught": "2026-09-26",
        "status": "Ongoing",
        "percent_completed": 50.0,
        "percent_covered": 50.0,
        "classwork": "Guided coding",
        "homework": "",  # Empty / not evidenced
    }
    xlsx_bytes = excel_service.generate_progression_workbook(rec)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[PRIMARY_SHEET_NAME]

    assert ws.cell(row=2, column=6).value in ("", None)
    assert ws.cell(row=2, column=5).value == "Guided coding"


def test_invalid_data(client):
    """Test 5 — Invalid data: Malformed progression produces a clear 422 validation error."""
    # Empty body
    res_empty = client.post("/progress/export", json={})
    assert res_empty.status_code == 422

    # Malformed object without progression fields
    res_malformed = client.post("/progress/export", json={"foo": "bar"})
    assert res_malformed.status_code == 422


def test_excel_formatting(sample_progression_payload):
    """Test 6 — Excel formatting: Checks bold header, freeze panes, autofilter, and widths."""
    xlsx_bytes = excel_service.generate_progression_workbook(sample_progression_payload)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[PRIMARY_SHEET_NAME]

    # Header bold
    assert ws.cell(row=1, column=1).font.bold is True
    # Freeze panes on A2
    assert ws.freeze_panes == "A2"
    # Auto-filter reference
    assert ws.auto_filter.ref is not None
    assert "A1:F" in ws.auto_filter.ref
    # Column widths configured
    assert ws.column_dimensions["A"].width > 10
    assert ws.column_dimensions["E"].width >= 30


def test_end_to_end_pipeline(client):
    """
    Test 7 — End-to-end integration:
    Curriculum -> Day 3 Analysis -> Day 4 Progress Calculation -> Day 5 Excel Export
    """
    # 1. Day 1 Curriculum
    curriculum = Curriculum(
        course_id="CS-101",
        course_name="Data Structures",
        lectures=[
            LecturePlan(
                week=1,
                day=1,
                lecture="LEC-1",
                content=["Arrays", "Array Traversal", "Time Complexity"],
                classwork=["Array Practice"],
                homework=["Homework 1"],
            )
        ],
    )

    # 2. Day 3 Mock Analysis
    analysis = LectureAnalysisResult(
        analysis_status="SUCCESS",
        matched_lecture=MatchedLecture(lecture_name="LEC-1", confidence=0.95),
        topics_taught=[
            TopicTaught(
                curriculum_topic="Arrays",
                evidence="Instructor discussed memory structure of static arrays.",
                confidence=0.95,
            ),
            TopicTaught(
                curriculum_topic="Array Traversal",
                evidence="Demonstrated for loop traversal across elements.",
                confidence=0.92,
            ),
        ],
        topics_not_evidenced=["Time Complexity"],
        classwork=[
            ClassworkPerformed(
                description="Array Practice",
                evidence="Students wrote for loop traversal code.",
                confidence=0.90,
            )
        ],
        homework=[
            HomeworkAssigned(
                description="Homework 1",
                evidence="Assigned practice questions 1 through 5.",
                confidence=0.90,
            )
        ],
    )

    # 3. Day 4 Progress Calculation
    progression = progress_engine.calculate_progress(
        analysis=analysis,
        curriculum=curriculum,
        lecture_code="LEC-1",
        date_taught="2026-09-26",
    )
    assert progression.lecture == "LEC-1"
    assert progression.percent_covered == 66.67
    assert progression.percent_completed == 66.67
    assert progression.status == "ONGOING"

    # 4. Day 5 Excel Export via API
    response = client.post(
        "/progress/export",
        json={
            "progression": progression.model_dump(),
            "file_name": "test_e2e_export.xlsx",
        },
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment; filename=" in response.headers["content-disposition"]

    # Verify binary workbook content
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    assert PRIMARY_SHEET_NAME in wb.sheetnames
    ws = wb[PRIMARY_SHEET_NAME]
    assert ws.cell(row=2, column=1).value == "2026-09-26"
    assert ws.cell(row=2, column=2).value == "Ongoing"
    assert float(ws.cell(row=2, column=3).value) == 66.67
    assert float(ws.cell(row=2, column=4).value) == 66.67
    assert "Array Practice" in ws.cell(row=2, column=5).value
    assert "Homework 1" in ws.cell(row=2, column=6).value


def test_api_export_update(client, sample_progression_payload):
    """Test 8 — API endpoint POST /progress/export/update multipart upload."""
    # Generate initial workbook
    initial_bytes = excel_service.generate_progression_workbook(sample_progression_payload)

    # Update with new lecture
    new_lecture = {
        "lecture": "LEC-2",
        "date_taught": "2026-09-28",
        "status": "Completed",
        "percent_completed": 100.0,
        "percent_covered": 100.0,
        "classwork": "Linked Lists Implementation",
        "homework": "Solve 3 LL problems",
    }

    import json
    files = {
        "file": ("existing.xlsx", initial_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    }
    data = {
        "progression_json": json.dumps(new_lecture)
    }

    response = client.post("/progress/export/update", files=files, data=data)
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb[PRIMARY_SHEET_NAME]
    assert ws.max_row == 3
    assert ws.cell(row=2, column=1).value == "2026-09-26"
    assert ws.cell(row=3, column=1).value == "2026-09-28"
    assert ws.cell(row=3, column=2).value == "Completed"


def test_download_endpoint(client):
    """Test 9 — API endpoint GET /progress/download returns course_progression.xlsx file."""
    response = client.get("/progress/download")
    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="course_progression.xlsx"' in response.headers["content-disposition"]
    assert len(response.content) > 0
