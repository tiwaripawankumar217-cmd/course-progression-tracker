import io
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
import openpyxl

from backend.main import app
from backend.services.progress_engine import progress_engine
from backend.models.progress import (
    TopicProgress,
    TopicStatus,
    LectureProgression,
    ClassworkItem,
    HomeworkItem,
)
from backend.services.excel_service import PRIMARY_SHEET_NAME


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def sample_review_progression():
    return LectureProgression(
        lecture="LEC-1",
        date_taught="2026-09-26",
        status="ONGOING",
        percent_completed=66.67,
        percent_covered=100.0,
        topics=[
            TopicProgress(name="Arrays", status=TopicStatus.COVERED, confidence=0.95, evidence="Contiguous array memory."),
            TopicProgress(name="Array Traversal", status=TopicStatus.COVERED, confidence=0.92, evidence="For loop scan."),
            TopicProgress(name="Time Complexity", status=TopicStatus.UNCERTAIN, confidence=0.45, evidence="Brief mention."),
        ],
        classwork=[ClassworkItem(description="Array Scan Coding")],
        homework=[HomeworkItem(description="Practice Set 1")],
    )


def test_recalculate_from_topics_override(sample_review_progression):
    """Test 1: Recalculates metrics deterministically when teacher marks uncertain topic as covered."""
    # Teacher marks Time Complexity as COVERED
    overridden_topics = [
        TopicProgress(name="Arrays", status=TopicStatus.COVERED, confidence=0.95),
        TopicProgress(name="Array Traversal", status=TopicStatus.COVERED, confidence=0.92),
        TopicProgress(name="Time Complexity", status=TopicStatus.COVERED, confidence=0.90),
    ]

    cov, comp, status_val = progress_engine.recalculate_from_topics(overridden_topics)
    assert cov == 100.0
    assert comp == 100.0
    assert status_val == "COMPLETED"


def test_recalculate_from_topics_uncover(sample_review_progression):
    """Test 2: Recalculates metrics deterministically when teacher marks a topic as NOT_COVERED."""
    # Teacher marks Time Complexity as NOT_COVERED
    overridden_topics = [
        TopicProgress(name="Arrays", status=TopicStatus.COVERED, confidence=0.95),
        TopicProgress(name="Array Traversal", status=TopicStatus.COVERED, confidence=0.92),
        TopicProgress(name="Time Complexity", status=TopicStatus.NOT_COVERED, confidence=None),
    ]

    cov, comp, status_val = progress_engine.recalculate_from_topics(overridden_topics)
    assert cov == 66.67
    assert comp == 66.67
    assert status_val == "ONGOING"


def test_approve_teacher_review_success(client, sample_review_progression, tmp_path, monkeypatch):
    """Test 3: POST /progress/review/approve records approval and signs off progression."""
    # Redirect excel write to temp directory for clean test isolation
    test_excel_path = tmp_path / "course_progression.xlsx"
    monkeypatch.chdir(tmp_path)

    payload = {
        "progression": sample_review_progression.model_dump(),
        "instructor_name": "Prof. Pawan Sharma",
        "comments": "Reviewed in lab; confirmed concepts understood.",
        "commit_to_excel": True,
    }

    response = client.post("/progress/review/approve", json=payload)
    assert response.status_code == 200
    data = response.json()

    assert data["success"] is True
    assert data["instructor_name"] == "Prof. Pawan Sharma"
    assert "approved successfully" in data["message"]
    assert data["excel_updated"] is True

    # Check warnings contain sign-off stamp
    warnings = data["approved_progression"]["warnings"]
    assert any("Approved by Instructor: Prof. Pawan Sharma" in w for w in warnings)
    assert any("Instructor Note: Reviewed in lab;" in w for w in warnings)

    # Verify that Excel file was created and contains lecture
    assert test_excel_path.exists()
    wb = openpyxl.load_workbook(test_excel_path)
    assert PRIMARY_SHEET_NAME in wb.sheetnames
    ws = wb[PRIMARY_SHEET_NAME]
    assert ws.cell(row=2, column=1).value == "2026-09-26"
    assert ws.cell(row=2, column=2).value == "Ongoing"


def test_approve_teacher_review_missing_name(client, sample_review_progression):
    """Test 4: Rejects approval when instructor name is missing or empty."""
    payload = {
        "progression": sample_review_progression.model_dump(),
        "instructor_name": "   ",  # Blank name
        "commit_to_excel": False,
    }

    response = client.post("/progress/review/approve", json=payload)
    assert response.status_code == 422
    assert "Instructor name is required" in response.json()["detail"]
